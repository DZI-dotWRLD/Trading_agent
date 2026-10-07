"""metrics.json: the small extract the dashboard and notifications read (ADR-0010).

Written next to each published output as ``<output stem>.metrics.json``.
Built from the validated, recalculated output (source tabs + generated tabs)
using the same definitions the validator checks, so it cannot drift from the
workbook. All money in EUR thousands.
"""
from __future__ import annotations

import json
from datetime import date, datetime
from pathlib import Path

import openpyxl

from agent.checks.metrics import headline
from agent.checks.reference import cash_schedule, monthly_summary, weekly_summary
from agent.checks.validate import company_name, resolve_assumptions
from agent.claude.memory.prior import prior_week
from agent.workbook.buckets import build_buckets
from agent.workbook.reader import TradingUpdate, WorkbookError, find_row

SCHEMA_VERSION = 1
SEGMENTS = {"cds": "Comparator Drug Sourcing", "cs": "Clinical Services (CS)", "ea": "Early Access (EA)",
            "total": "Total"}


def metrics_path(output: Path) -> Path:
    return output.with_name(f"{output.stem}.metrics.json")


def build_metrics(output: Path, context: dict, reports_dir: Path | None, company: str, revision: int = 1) -> dict:
    tu = TradingUpdate(output, context.get("source_map"))  # the validated output carries the source tabs
    year, cw = tu.iso_week
    prior = prior_week(tu, reports_dir)
    vals = openpyxl.load_workbook(output, data_only=True)
    sched = cash_schedule(tu.orderbook(), build_buckets(tu.report_date, int(context["assumptions"].get("forward_weeks", 28))),
                          resolve_assumptions(context, tu, vals))
    name = company_name(context)
    labels_m, monthly = monthly_summary(sched, name)
    return {
        "schema": SCHEMA_VERSION,
        "company": company,
        "cw": cw, "year": year, "revision": revision,
        "report_date": tu.report_date.isoformat(), "cash_date": tu.cash_date.isoformat(),
        "source_file": context["source_file"], "output_file": output.name,
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "headline": headline(tu),
        "prior": {"source": prior.source, "file": prior.path.name if prior.path else None,
                  "report_date": prior.report_date.isoformat() if isinstance(prior.report_date, date) else None,
                  "values": prior.values},
        "segments": _segments(tu),
        "customers": _customers(tu, prior.customers),
        "cash": _cash(tu),
        "gp_history": _gp_history(tu, cw),
        "cash_summary": {
            "weekly": {"labels": [b.label for b in sched.buckets],
                       "week_ending": [b.week_ending.isoformat() if b.week_ending else None for b in sched.buckets],
                       "rows": weekly_summary(sched, name)},
            "monthly": {"labels": labels_m, "rows": monthly},
        },
    }


def write_metrics(output: Path, data: dict) -> Path:
    p = metrics_path(output)
    tmp = p.with_name(f".{p.name}.partial")
    tmp.write_text(json.dumps(data, indent=2, default=str), encoding="utf-8")
    tmp.replace(p)
    return p


def _segments(tu: TradingUpdate) -> dict:
    """Segment rows through the source map. The per-segment revenue / GP vs prior year and budget come from the
    template's Summary layout; for another layout they are left empty (the dashboard shows n/a)."""
    empty = {"ytd": None, "py": None, "budget": None}
    template = tu.layout.summary_sheet == "Summary" and tu.layout.summary_value_col == "F"
    sm = tu.wb[tu.layout.summary_sheet] if template else None

    def summary_row(label: str, start: int | None) -> dict:
        if sm is None or start is None:
            return dict(empty)
        try:
            r = find_row(sm, "C", label, start, start + 10)
        except WorkbookError:
            return dict(empty)
        num = lambda c: float(sm[f"{c}{r}"].value) if isinstance(sm[f"{c}{r}"].value, (int, float)) else None  # noqa: E731
        return {"ytd": num("F"), "py": num("G"), "budget": num("J")}

    def top(label: str) -> int | None:
        try:
            return find_row(sm, "C", label, 1, 40) if sm is not None else None
        except WorkbookError:
            return None

    rev_top, gp_top = top("Adjusted Revenue**"), top("Gross Profit")
    out = {}
    for key in SEGMENTS:
        out[key] = tu.segment(key) or {f: None for f in ("inv_rev", "inv_gp", "ob_rev", "ob_gp", "iob_rev", "iob_gp")}
    rev_labels = {"cds": "Comparator Drug Sourcing", "cs": "Clinical Services (CS)",
                  "ea": "Early Access (EA) Program Fees", "total": "Total Adjusted Revenue**"}
    gp_labels = {"cds": "Comparator Drug Sourcing", "cs": "CS", "ea": "EA", "total": "Total Gross Profit"}
    for key in SEGMENTS:
        out[key]["adj_revenue"] = summary_row(rev_labels[key], rev_top)
        out[key]["gp"] = summary_row(gp_labels[key], gp_top)
    out["reported_revenue"] = summary_row("Total Reported Revenue", rev_top)
    if not template:  # the totals and budgets the source map does locate
        out["total"]["adj_revenue"].update(ytd=tu.summary("adj_revenue"), budget=tu.summary("adj_revenue", budget=True))
        out["total"]["gp"].update(ytd=tu.summary("total_gp"), budget=tu.summary("total_gp", budget=True))
    return out


def _customers(tu: TradingUpdate, prior: dict[str, float]) -> list[dict]:
    rows = []
    for name, v in tu.cds_customers().items():
        rows.append({"name": name, "segment": "CDS", "iob_gp": v["iob_gp"], "iob_rev": v["iob_rev"],
                     "inv_gp": v["inv_gp"], "ob_gp": v["ob_gp"],
                     "prior_iob_gp": prior.get(name) if prior else None})
    return sorted(rows, key=lambda r: -r["iob_gp"])


def _cash(tu: TradingUpdate) -> dict:
    """Bank accounts (the template's Cash tab only) and the balance sheet, this week and last."""
    accounts = []
    template = tu.layout.sheets.get("cash") == "Cash" and tu.layout.opening_bank_cash is not None
    ws = tu.wb["Cash"] if template else None
    for r in (range(8, 29) if template else ()):
        eur = ws[f"I{r}"].value
        if isinstance(eur, (int, float)):
            prior = ws[f"L{r}"].value
            accounts.append({"affiliate": (ws[f"C{r}"].value or "").strip(), "bank": (ws[f"F{r}"].value or "").strip(),
                             "country": ws[f"E{r}"].value, "eur_k": eur / 1000,
                             "prior_eur_k": prior / 1000 if isinstance(prior, (int, float)) else None})
    bs, bs_prior = tu.balance_sheet("current"), tu.balance_sheet("prior")
    return {"accounts": sorted(accounts, key=lambda a: -a["eur_k"]), "balance_sheet": bs, "balance_sheet_prior": bs_prior}


def _gp_history(tu: TradingUpdate, cw: int) -> list[dict]:
    hist = []
    for w in range(2, cw + 1):
        h = tu.weekly_gp(w)
        if h and any(h[k] for k in ("inv_gp_total", "ob_gp_total")):
            hist.append({"cw": w, "invoiced": h["inv_gp_total"], "orderbook": h["ob_gp_total"],
                         "total": h["inv_gp_total"] + h["ob_gp_total"]})
    return hist
