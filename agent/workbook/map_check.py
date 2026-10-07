"""Does a source map really fit this workbook? (ADR-0024)

``resolve()`` only proves that every mapped label and header exists. A *wrong* map (sales and purchase
swapped, the CDS label pointing at the CS row, cash pointing at the debt row) still resolves. So
``check_map`` also checks value types and accounting identities that hold in every real file: a swapped
or misplaced column or row breaks at least one of them.

``check_map(wb_values, source_map)`` takes a workbook opened with ``data_only=True`` and returns
``MapCheck`` items. Severity "error" must pass for the map to be used (``passed(checks)``); "warning" is
plausibility shown to the reviewer.

Identities kept, verified on ``samples/Trading_update_CW36_2026.xlsx``, the generated weeks CW33–40
(``data/generated``, also the Meridian variant) and the six test companies' CW37–39 files
(``data/fulltest/files/*``), 28 files. On all of them every identity below holds to float noise
(largest difference about 4e-12), so the tolerances are generous for structure, not for rounding:

- orderbook, per line: ``gp = sales - purchase``. A line agrees within max(1, 0.5 % of |sales|) EUR;
  the check passes when at least 98 % of lines agree (real files: 100 %).
- segments, per mapped row: ``iob_rev = inv_rev + ob_rev`` and ``iob_gp = inv_gp + ob_gp``.
- segments: cds + cs + ea = total in all six columns, when all three rows are mapped. When a segment is
  ``null`` the identity can't be checked exactly; then the mapped segments' revenue must not exceed the
  total's (``|sum| <= |total|``, a warning only, since no real file without EA was available to verify it).
- CDS customers: the customer rows sum to the block's total row, and that total row equals the CDS segment
  row, in all six segment columns.
- balance sheet, in the current column and (if mapped) the prior column:
  ``net_debt = debt - cash``; ``net_debt_incl_pharma = debt - cash_incl_pharma``;
  ``cash_incl_pharma = cash + holdco``.
- summary ``total_gp`` = the segments' total ``iob_gp`` (the Summary's GP is invoiced + orderbook).
- CDS GP <= CDS revenue (margin at most 100 %) for inv, ob and iob on the CDS segment row and every
  CDS-customer row: ``gp <= |rev| + tolerance``. Not on the CS row: Meridian Health and Northbridge
  Clinical (CW37–39) book CS GP above CS revenue (e.g. inv GP 798 vs revenue 750), so the candidate
  "on every segment row" was narrowed to CDS. ``iob = inv + ob`` and the sums are symmetric in ``inv_*``/``ob_*``, so this
  is what catches swapped invoiced and orderbook GP columns (a customer with no open orders would show
  orderbook GP without orderbook revenue).
- weekly GP history (when mapped): this week's column (header ``CWnn`` = the report date's ISO week or
  Excel ``WEEKNUM(d, 1)``) holds the segments' inv / ob / iob GP in its three blocks (CDS, CS, EA, total).
  Older columns are not compared: the sample's history doesn't add up (``Weekly GP!AL143:AL145``, see
  CLAUDE.md gotchas). Skipped (a passing warning) when this week's column is absent or a segment is null.

Identity tolerance (segments, customers, balance sheet, summary, weekly GP):
``|a - b| <= 0.5 + 1e-4 * max(|a|, |b|)`` in the sheet's own units (EUR 000s in the template).

Candidates dropped:
- Summary ``adj_revenue`` vs the segments' ``iob_rev``: not equal in any file (adjusted revenue is a
  different measure: 300,825 vs 369,995 in the sample), so no identity links them.
- Orderbook totals vs the CDS segment's orderbook columns (sum of sales / 1000 vs ``ob_rev``, gp vs
  ``ob_gp``): close in the sample (100,104 vs 100,648) but not equal, and far apart in generated weeks
  (CW33: 126,276 vs 75,004; gp 16,958 vs 6,980). The orderbook isn't the same snapshot as the P&L.
- CDS customers sorted by invoiced GP (the block's heading says so): the sample's rows are not sorted
  (3 of 13 adjacent pairs ascend in ``inv_gp``), so order can't tell columns apart.

Type checks (errors): report date (and cash date, if mapped) are dates; opening bank cash is a number;
orderbook ``sales``/``purchase``/``gp`` numeric on at least 98 % of lines; ``so_date`` a date on at least
95 % of lines (real files: 100 %; a swap with ``po_date`` brings in the "PM fees" text, about 17 %);
``po_date`` a date or the text "PM fees" on at least 95 %; SO numbers non-empty on lines that carry a
sales value; segment and balance-sheet value cells numeric.

Plausibility (warning): on lines with both dates, the supplier is paid on or before the customer pays on
most lines (real files: 91–99 % of lines; with the two date columns swapped it would be 1–9 %).
The check warns when fewer than half do.
"""
from __future__ import annotations

import re
from dataclasses import asdict, dataclass
from datetime import date, datetime

from openpyxl.utils import column_index_from_string, get_column_letter

from agent.workbook.source_map import Layout, MapError, norm, resolve

ABS_TOL = 0.5
REL_TOL = 1e-4
NUMERIC_SHARE = 0.98
DATE_SHARE = 0.95
GP_LINE_SHARE = 0.98
PO_BEFORE_SO_SHARE = 0.5


@dataclass
class MapCheck:
    name: str
    passed: bool
    detail: str = ""
    severity: str = "error"  # "error" | "warning"

    def to_dict(self) -> dict:
        return asdict(self)


def passed(checks: list[MapCheck]) -> bool:
    """True when every error-severity check passes (warnings are for the reviewer)."""
    return all(c.passed for c in checks if c.severity == "error")


def failures(checks: list[MapCheck], severity: str | None = None) -> list[MapCheck]:
    return [c for c in checks if not c.passed and (severity is None or c.severity == severity)]


def _num(v) -> float | None:
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def _is_date(v) -> bool:
    return isinstance(v, (datetime, date))


def _close(a: float, b: float) -> bool:
    return abs(a - b) <= ABS_TOL + REL_TOL * max(abs(a), abs(b))


def _fmt(x: float | None) -> str:
    return "n/a" if x is None else f"{x:,.2f}"


def check_map(wb_values, source_map: dict) -> list[MapCheck]:
    """All checks of ``source_map`` against ``wb_values`` (opened with data_only=True)."""
    try:
        lay = resolve(wb_values, source_map)
    except MapError as e:
        return [MapCheck("map resolves", False, "; ".join(e.problems))]
    checks = [MapCheck("map resolves", True, "every required input found")]
    if lay.problems:
        checks.append(MapCheck("optional inputs found", False, "; ".join(lay.problems), "warning"))
    checks += _check_dates(wb_values, lay)
    checks += _check_orderbook(wb_values, lay)
    checks += _check_segments(wb_values, lay)
    checks += _check_balance_sheet(wb_values, lay)
    checks += _check_summary(wb_values, lay)
    checks += _check_weekly_gp(wb_values, lay)
    return checks


# --------------------------------------------------------------------------- dates
def _check_dates(wb, lay: Layout) -> list[MapCheck]:
    out = []
    for name, ref in (("report date", lay.report_date), ("cash date", lay.cash_date)):
        if ref is None:
            continue
        v = wb[ref[0]][ref[1]].value
        out.append(MapCheck(f"{name} is a date", _is_date(v), f"'{ref[0]}'!{ref[1]} holds {v!r}"))
    if lay.opening_bank_cash:
        s, c = lay.opening_bank_cash
        v = wb[s][c].value
        out.append(MapCheck("opening bank cash is a number", _num(v) is not None, f"'{s}'!{c} holds {v!r}"))
    return out


# --------------------------------------------------------------------------- orderbook
def _orderbook_rows(wb, lay: Layout) -> list[dict]:
    ws = wb[lay.ob_sheet]
    cols = {f: c for f, c in lay.ob_cols.items() if c}
    rows = []
    for r in range(lay.ob_first_row, lay.ob_last_row + 1):
        vals = {f: ws[f"{c}{r}"].value for f, c in cols.items()}
        if all(v in (None, "") for v in vals.values()):
            continue
        vals["_row"] = r
        rows.append(vals)
    return rows


def _share(n: int, total: int) -> float:
    return n / total if total else 0.0


def _check_orderbook(wb, lay: Layout) -> list[MapCheck]:
    rows = _orderbook_rows(wb, lay)
    lines = [x for x in rows if x.get("so") not in (None, "")]
    n = len(lines)
    where = f"'{lay.ob_sheet}' rows {lay.ob_first_row}-{lay.ob_last_row}"
    out = []
    if n == 0:
        return [MapCheck("orderbook has lines", False, f"no SO numbers in column {lay.ob_cols['so']} of {where}")]

    for f in ("sales", "purchase", "gp"):
        bad = [x["_row"] for x in lines if _num(x.get(f)) is None]
        out.append(MapCheck(f"orderbook {f} is numeric", _share(n - len(bad), n) >= NUMERIC_SHARE,
                            f"{n - len(bad)} of {n} lines numeric in column {lay.ob_cols[f]}"
                            + (f"; not numeric at rows {bad[:5]}" if bad else "")))

    bad = [x["_row"] for x in lines if not _is_date(x.get("so_date"))]
    out.append(MapCheck("orderbook so_date holds dates", _share(n - len(bad), n) >= DATE_SHARE,
                        f"{n - len(bad)} of {n} lines are dates in column {lay.ob_cols['so_date']}"
                        + (f"; not a date at rows {bad[:5]}" if bad else "")))

    bad = [x["_row"] for x in lines if not (_is_date(x.get("po_date")) or norm(x.get("po_date")) == "pm fees")]
    out.append(MapCheck("orderbook po_date holds dates or 'PM fees'", _share(n - len(bad), n) >= DATE_SHARE,
                        f"{n - len(bad)} of {n} lines in column {lay.ob_cols['po_date']}"
                        + (f"; other values at rows {bad[:5]}" if bad else "")))

    with_sales = [x for x in rows if _num(x.get("sales"))]
    no_so = [x["_row"] for x in with_sales if x.get("so") in (None, "")]
    out.append(MapCheck("orderbook SO numbers are non-empty", _share(len(with_sales) - len(no_so), len(with_sales))
                        >= DATE_SHARE, f"{len(no_so)} of {len(with_sales)} lines with a sales value have no SO number"
                        + (f" (rows {no_so[:5]})" if no_so else "")))

    # gp = sales - purchase, line by line
    num_lines = [x for x in lines if None not in (_num(x.get("sales")), _num(x.get("purchase")), _num(x.get("gp")))]
    off = []
    for x in num_lines:
        s, p, g = _num(x["sales"]), _num(x["purchase"]), _num(x["gp"])
        if abs(g - (s - p)) > max(1.0, 0.005 * abs(s)):
            off.append(x["_row"])
    m = len(num_lines)
    out.append(MapCheck("orderbook gp = sales - purchase", m > 0 and _share(m - len(off), m) >= GP_LINE_SHARE,
                        f"{m - len(off)} of {m} lines agree"
                        + (f"; disagree at rows {off[:5]} (columns sales {lay.ob_cols['sales']}, purchase "
                           f"{lay.ob_cols['purchase']}, gp {lay.ob_cols['gp']})" if off else "")))

    # plausibility: suppliers are paid on or before customers pay
    dated = [x for x in lines if _is_date(x.get("so_date")) and _is_date(x.get("po_date"))]
    before = sum(1 for x in dated if _as_dt(x["po_date"]) <= _as_dt(x["so_date"]))
    out.append(MapCheck("po_date mostly on or before so_date", bool(dated) and _share(before, len(dated)) >= PO_BEFORE_SO_SHARE,
                        f"{before} of {len(dated)} dated lines have po_date <= so_date", "warning"))
    return out


def _as_dt(v) -> datetime:
    return v if isinstance(v, datetime) else datetime(v.year, v.month, v.day)


# --------------------------------------------------------------------------- segments and customers
def _check_segments(wb, lay: Layout) -> list[MapCheck]:
    ws = wb[lay.seg_sheet]
    cols = lay.seg_cols
    out = []
    segs = {k: r for k, r in lay.seg_rows.items() if r}

    bad = [f"{cols[c]}{r}={ws[f'{cols[c]}{r}'].value!r}" for r in segs.values() for c in cols
           if _num(ws[f"{cols[c]}{r}"].value) is None]
    out.append(MapCheck("segment values are numeric", not bad,
                        "all numeric" if not bad else f"'{lay.seg_sheet}' not numeric: {', '.join(bad[:6])}"))
    if bad:
        return out

    def v(r: int, c: str) -> float:
        return _num(ws[f"{cols[c]}{r}"].value) or 0.0

    off = []
    for k, r in segs.items():
        for kind in ("rev", "gp"):
            a, b = v(r, f"iob_{kind}"), v(r, f"inv_{kind}") + v(r, f"ob_{kind}")
            if not _close(a, b):
                off.append(f"{k} row {r} {kind}: iob {_fmt(a)} vs inv+ob {_fmt(b)}")
    out.append(MapCheck("segments: iob = inv + ob", not off, "; ".join(off) or "holds on every segment row"))

    parts = [s for s in ("cds", "cs", "ea") if s in segs]
    tot = segs["total"]
    if len(parts) == 3:
        off = [f"{c} ({cols[c]}): segments {_fmt(sum(v(segs[s], c) for s in parts))} vs total {_fmt(v(tot, c))}"
               for c in cols if not _close(sum(v(segs[s], c) for s in parts), v(tot, c))]
        out.append(MapCheck("segments sum to the total row", not off, "; ".join(off) or "all six columns tie"))
    else:
        off = [f"{c}: {_fmt(sum(v(segs[s], c) for s in parts))} > total {_fmt(v(tot, c))}"
               for c in ("inv_rev", "ob_rev", "iob_rev")
               if abs(sum(v(segs[s], c) for s in parts)) > abs(v(tot, c)) + ABS_TOL]
        out.append(MapCheck("mapped segments don't exceed the total", not off,
                            "; ".join(off) or f"only {parts} mapped; their revenue is within the total", "warning"))
    if len(set(segs.values())) < len(segs):
        out.append(MapCheck("segment rows are distinct", False, f"two segments map to the same row: {segs}"))

    # customers: the rows sum to the block's total row, which equals the CDS segment row
    first, end = lay.cust_first_row, lay.cust_end_row
    if first is not None and end is not None:
        off = []
        for c in cols:
            s = sum(_num(ws[f"{cols[c]}{r}"].value) or 0.0 for r in range(first, end))
            if not _close(s, v(end, c)):
                off.append(f"{c} ({cols[c]}): rows {first}-{end - 1} sum {_fmt(s)} vs row {end} {_fmt(v(end, c))}")
        out.append(MapCheck("CDS customers sum to their total row", not off and end > first,
                            "; ".join(off) or (f"rows {first}-{end - 1} tie in all six columns" if end > first
                                               else "the block has no customer rows")))
        off = [f"{c} ({cols[c]}): customers' total {_fmt(v(end, c))} vs CDS segment {_fmt(v(segs['cds'], c))}"
               for c in cols if not _close(v(end, c), v(segs["cds"], c))]
        out.append(MapCheck("CDS customers' total = CDS segment row", not off, "; ".join(off) or "all six columns tie"))

    # CDS GP never exceeds CDS revenue (margin <= 100 %), on the CDS row and its customers: tells inv_gp from
    # ob_gp. Not on CS: service revenue can carry GP above the revenue booked (seen in two test companies).
    rows = [("cds", segs["cds"])]
    if first is not None and end is not None:
        rows += [(f"customer row {r}", r) for r in range(first, end)]
    off = []
    for k, r in rows:
        for kind in ("inv", "ob", "iob"):
            rev, gp = _num(ws[f"{cols[kind + '_rev']}{r}"].value), _num(ws[f"{cols[kind + '_gp']}{r}"].value)
            if rev is not None and gp is not None and gp > abs(rev) + ABS_TOL + REL_TOL * abs(rev):
                off.append(f"{k} ({r}) {kind}: gp {_fmt(gp)} ({cols[kind + '_gp']}) > revenue {_fmt(rev)} "
                           f"({cols[kind + '_rev']})")
    out.append(MapCheck("CDS GP <= revenue on the CDS row and its customers", not off,
                        "; ".join(off[:5]) + (f" (+{len(off) - 5} more)" if len(off) > 5 else "")
                        or "holds on every row"))
    return out


# --------------------------------------------------------------------------- balance sheet
def _check_balance_sheet(wb, lay: Layout) -> list[MapCheck]:
    ws = wb[lay.seg_sheet]
    rows = {k: r for k, r in lay.bs_rows.items() if r}
    out = []
    for col, which in ((lay.bs_current_col, "current"), (lay.bs_prior_col, "prior")):
        if not col:
            continue
        vals = {k: ws[f"{col}{r}"].value for k, r in rows.items()}
        bad = [f"{k} {col}{rows[k]}={x!r}" for k, x in vals.items() if _num(x) is None]
        out.append(MapCheck(f"balance sheet ({which}) values are numeric", not bad,
                            "all numeric" if not bad else "not numeric: " + ", ".join(bad)))
        if bad:
            continue
        b = {k: _num(x) for k, x in vals.items()}
        idents = [("net_debt = debt - cash", "net_debt", ("debt", "cash")),
                  ("net_debt_incl_pharma = debt - cash_incl_pharma", "net_debt_incl_pharma", ("debt", "cash_incl_pharma"))]
        for name, lhs, (pos, neg) in idents:
            if lhs in b and pos in b and neg in b:
                ok = _close(b[lhs], b[pos] - b[neg])
                out.append(MapCheck(f"balance sheet ({which}): {name}", ok,
                                    f"{lhs} {_fmt(b[lhs])} vs {pos} - {neg} = {_fmt(b[pos] - b[neg])} (column {col})"))
        if {"cash_incl_pharma", "cash", "holdco"} <= set(b):
            ok = _close(b["cash_incl_pharma"], b["cash"] + b["holdco"])
            out.append(MapCheck(f"balance sheet ({which}): cash_incl_pharma = cash + holdco", ok,
                                f"{_fmt(b['cash_incl_pharma'])} vs {_fmt(b['cash'] + b['holdco'])} (column {col})"))
    if len(set(rows.values())) < len(rows):
        out.append(MapCheck("balance-sheet rows are distinct", False, f"two inputs map to the same row: {rows}"))
    return out


# --------------------------------------------------------------------------- summary and weekly GP
def _check_summary(wb, lay: Layout) -> list[MapCheck]:
    if not lay.summary_sheet or not lay.summary_value_col:
        return []
    ws = wb[lay.summary_sheet]
    out = []
    for k, r in lay.summary_rows.items():
        if r is None:
            continue
        x = ws[f"{lay.summary_value_col}{r}"].value
        out.append(MapCheck(f"summary {k} is numeric", _num(x) is not None,
                            f"'{lay.summary_sheet}'!{lay.summary_value_col}{r} holds {x!r}"))
    r = lay.summary_rows.get("total_gp")
    seg_ws, tot = wb[lay.seg_sheet], lay.seg_rows.get("total")
    a = _num(ws[f"{lay.summary_value_col}{r}"].value) if r else None
    b = _num(seg_ws[f"{lay.seg_cols['iob_gp']}{tot}"].value) if tot else None
    if a is not None and b is not None:
        out.append(MapCheck("summary total_gp = segments' total iob_gp", _close(a, b),
                            f"'{lay.summary_sheet}'!{lay.summary_value_col}{r} {_fmt(a)} vs "
                            f"'{lay.seg_sheet}'!{lay.seg_cols['iob_gp']}{tot} {_fmt(b)}"))
    return out


def _week_labels(d) -> set[str]:
    """CW labels the report week may carry: ISO week, or Excel's WEEKNUM(d, 1) (weeks start on Sunday)."""
    d = d.date() if isinstance(d, datetime) else d
    jan1 = date(d.year, 1, 1)
    weeknum = ((d - jan1).days + (jan1.isoweekday() % 7)) // 7 + 1
    return {f"cw{d.isocalendar()[1]:02d}", f"cw{weeknum:02d}"}


def _check_weekly_gp(wb, lay: Layout) -> list[MapCheck]:
    if not lay.wgp_sheet:
        return []
    ws = wb[lay.wgp_sheet]
    first = column_index_from_string(lay.wgp_first_col or "A")
    hdr = {c: norm(ws.cell(lay.wgp_header_row, c).value) for c in range(first, ws.max_column + 1)}
    cws = {c: h for c, h in hdr.items() if re.fullmatch(r"cw\s?\d{1,2}", h)}
    out = [MapCheck("weekly GP header row has CW columns", bool(cws),
                    f"'{lay.wgp_sheet}' row {lay.wgp_header_row}: {len(cws)} CW labels from column {lay.wgp_first_col}")]
    rd = wb[lay.report_date[0]][lay.report_date[1]].value
    if not cws or not _is_date(rd):
        return out
    want = _week_labels(rd)
    col = next((c for c, h in cws.items() if h.replace(" ", "") in want), None)
    segs = lay.seg_rows
    if col is None or not all(segs.get(s) for s in ("cds", "cs", "ea", "total")):
        out.append(MapCheck("weekly GP: this week's column = segment GP", True,
                            "not checked (" + ("this week's CW column not in the history" if col is None
                                               else "a segment is not mapped") + ")", "warning"))
        return out
    sw = wb[lay.seg_sheet]
    off = []
    for block, top in (lay.wgp_blocks or {}).items():
        seg_col = lay.seg_cols.get(f"{block}_gp")
        for i, s in enumerate(("cds", "cs", "ea", "total")):
            a, b = _num(ws.cell(top + i, col).value), _num(sw[f"{seg_col}{segs[s]}"].value)
            if a is None or b is None or not _close(a, b):
                off.append(f"{block} {s}: {get_column_letter(col)}{top + i} {_fmt(a)} vs {seg_col}{segs[s]} {_fmt(b)}")
    out.append(MapCheck("weekly GP: this week's column = segment GP", not off,
                        "; ".join(off[:6]) or f"column {get_column_letter(col)} ties to inv/ob/iob GP of every segment"))
    return out
