"""Read a trading-update workbook through its source map (ADR-0024): by label and header, never by guesswork.

Every input is located by ``agent.workbook.source_map.resolve``: Inceptua's own layout by default, or a
company's approved map. Optional inputs a company's file doesn't have read as ``None``.

Everything here works on a workbook opened with ``data_only=True`` (cached values). Source files from
portfolio companies are saved by Excel/LibreOffice and carry cached values; generated or agent-written files
must be recalculated first (see ``agent.workbook.recalc``).
"""
from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

import openpyxl
from openpyxl.utils import column_index_from_string
from openpyxl.worksheet.worksheet import Worksheet

from agent.workbook.source_map import INCEPTUA_MAP, MapError, norm, resolve

# Kept for the test-data generators, which only ever write Inceptua's layout.
ORDERBOOK_HEADER_ROW = 7
CUSTOMER_LEVEL_COLS = INCEPTUA_MAP["segments"]["columns"]


class WorkbookError(Exception):
    pass


def _norm(v) -> str:
    return norm(v)


def as_date(v) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return None


@dataclass
class OrderLine:
    row: int
    so: str
    customer: str
    item: str
    sales: float
    purchase: float
    gp: float
    gp_pct: float | None
    so_date: datetime | None
    po_date: datetime | str | None  # datetime, or the text "PM fees"

    @property
    def is_pm_fee(self) -> bool:
        return isinstance(self.po_date, str)


def find_row(ws: Worksheet, col: str, label: str, start: int = 1, end: int | None = None) -> int:
    """First row in ``col`` whose text equals ``label`` (case/whitespace-insensitive)."""
    target = _norm(label)
    for r in range(start, (end or ws.max_row) + 1):
        if _norm(ws[f"{col}{r}"].value) == target:
            return r
    raise WorkbookError(f"{ws.title}: label {label!r} not found in column {col}")


class TradingUpdate:
    """Cached-value view of one trading-update workbook, read through its source map."""

    def __init__(self, path: str | Path, source_map: dict | None = None):
        self.path = Path(path)
        self.map = source_map or INCEPTUA_MAP
        self.wb = openpyxl.load_workbook(self.path, data_only=True)
        try:
            self.layout = resolve(self.wb, self.map)
        except MapError as e:
            raise WorkbookError(f"{self.path.name}: {e}") from e

    def _cell(self, ref: tuple[str, str] | None):
        return self.wb[ref[0]][ref[1]].value if ref else None

    # --- dates -----------------------------------------------------------
    @property
    def report_date(self) -> date:
        d = as_date(self._cell(self.layout.report_date))
        if d is None:
            raise WorkbookError(f"{self.layout.report_date[0]}!{self.layout.report_date[1]} (report date) is not a date")
        return d

    @property
    def cash_date(self) -> date:
        """The cash snapshot date; the report date when the file has none."""
        return as_date(self._cell(self.layout.cash_date)) or self.report_date

    @property
    def iso_week(self) -> tuple[int, int]:
        y, w, _ = self.report_date.isocalendar()
        return y, w

    @property
    def opening_bank_cash(self) -> float | None:
        """Total bank cash the cash schedule starts from (Cash!I29 in the template), or None if not in the file."""
        v = self._cell(self.layout.opening_bank_cash)
        return float(v) if isinstance(v, (int, float)) else None

    # --- orderbook -------------------------------------------------------
    def orderbook_columns(self) -> dict[str, int]:
        return {f: column_index_from_string(c) for f, c in self.layout.ob_cols.items() if c}

    def orderbook_last_row(self) -> int:
        return self.layout.ob_last_row

    def orderbook(self) -> list[OrderLine]:
        ws = self.wb[self.layout.ob_sheet]
        c = self.orderbook_columns()

        def get(r, f):
            return ws.cell(r, c[f]).value if f in c else None

        lines = []
        for r in range(self.layout.ob_first_row, self.layout.ob_last_row + 1):
            so = get(r, "so")
            if so in (None, ""):
                continue
            po, sd = get(r, "po_date"), get(r, "so_date")
            lines.append(OrderLine(
                row=r, so=str(so), customer=get(r, "customer") or "", item=get(r, "item") or "",
                sales=float(get(r, "sales") or 0), purchase=float(get(r, "purchase") or 0),
                gp=float(get(r, "gp") or 0), gp_pct=get(r, "gp_pct"),
                so_date=sd if isinstance(sd, datetime) else None,
                po_date=po if isinstance(po, (datetime, str)) else None,
            ))
        return lines

    # --- P&L -------------------------------------------------------------
    def summary(self, key: str, budget: bool = False) -> float | None:
        """Summary-tab total by meaning: ``adj_revenue`` or ``total_gp``; its full-year budget with ``budget``."""
        L = self.layout
        row, col = L.summary_rows.get(key), (L.summary_budget_col if budget else L.summary_value_col)
        if not (L.summary_sheet and row and col):
            return None
        return _num(self.wb[L.summary_sheet][f"{col}{row}"].value)

    def summary_value(self, label: str, col: str | None = None) -> float:
        """Legacy lookup by label (Inceptua's labels), used by the test-data generators."""
        L = self.layout
        key = next((k for k, lab in (self.map.get("summary") or {}).get("rows", {}).items() if lab and _norm(lab) == _norm(label)), None)
        budget = col is not None and col == L.summary_budget_col
        v = self.summary(key, budget) if key else None
        if v is None:
            raise WorkbookError(f"summary value {label!r} not in this file")
        return v

    def segment(self, key: str) -> dict[str, float] | None:
        """One segment row by meaning (``cds``, ``cs``, ``ea``, ``total``) -> {inv_rev .. iob_gp}, or None."""
        row = self.layout.seg_rows.get(key)
        return {f: self.customer_level(row, f) for f in self.layout.seg_cols} if row else None

    def segment_row(self, label: str) -> int:
        """Legacy: the row of a segment by its Inceptua label or its key."""
        rows = (self.map.get("segments") or {}).get("rows", {})
        key = label if label in self.layout.seg_rows else next(
            (k for k, lab in rows.items() if lab and _norm(lab) == _norm(label)), None)
        row = self.layout.seg_rows.get(key) if key else None
        if not row:
            raise WorkbookError(f"segment {label!r} not in this file")
        return row

    def customer_level(self, row: int, field: str) -> float:
        return _num(self.wb[self.layout.seg_sheet][f"{self.layout.seg_cols[field]}{row}"].value)

    def cds_customers(self) -> dict[str, dict[str, float]]:
        """CDS customer rows -> {field: value}. Row order changes weekly; key by name."""
        L = self.layout
        ws = self.wb[L.seg_sheet]
        out = {}
        for r in range(L.cust_first_row, L.cust_end_row):
            name = ws[f"{L.cust_label_col}{r}"].value
            if name:
                out[str(name)] = {f: self.customer_level(r, f) for f in L.seg_cols}
        return out

    def balance_sheet(self, which: str = "current") -> dict[str, float] | None:
        """Balance-sheet block: ``current`` (or "E") this week, ``prior`` (or "H") last week; None if absent.

        A file without a Pharma Model reports cash and net debt incl. Pharma equal to plain cash and net debt.
        """
        L = self.layout
        col = L.bs_prior_col if which in ("prior", "H") else L.bs_current_col
        if not col:
            return None
        ws = self.wb[L.seg_sheet]
        val = lambda k: _num(ws[f"{col}{L.bs_rows[k]}"].value) if L.bs_rows.get(k) else None  # noqa: E731
        out = {k: val(k) for k in ("cash", "debt", "net_debt", "cash_incl_pharma", "net_debt_incl_pharma")}
        out["cash_incl_pharma"] = out["cash"] if out["cash_incl_pharma"] is None else out["cash_incl_pharma"]
        out["net_debt_incl_pharma"] = (out["net_debt"] if out["net_debt_incl_pharma"] is None
                                       else out["net_debt_incl_pharma"])
        out["holdco_cash"] = val("holdco") or 0.0
        return out

    def balance_sheet_dates(self) -> tuple[date | None, date | None]:
        """(this week's, last week's) snapshot dates from the row below the balance-sheet heading."""
        L = self.layout
        ws = self.wb[L.seg_sheet]
        get = lambda col: as_date(ws[f"{col}{L.bs_date_row}"].value) if col and L.bs_date_row else None  # noqa: E731
        return get(L.bs_current_col), get(L.bs_prior_col)

    # --- Weekly GP history -------------------------------------------------
    def weekly_gp_column(self, cw: int, header_row: int | None = None) -> int | None:
        L = self.layout
        if not L.wgp_sheet:
            return None
        ws = self.wb[L.wgp_sheet]
        want = f"cw{cw:02d}"
        first = column_index_from_string(L.wgp_first_col or "E")
        for c in range(first, ws.max_column + 1):
            if _norm(ws.cell(header_row or L.wgp_header_row, c).value) == want:
                return c
        return None

    def weekly_gp(self, cw: int) -> dict[str, float] | None:
        """In-file history for one CW: invoiced / orderbook / I+OB GP by segment (EUR 000s), or None."""
        L = self.layout
        c = self.weekly_gp_column(cw)
        if c is None:
            return None
        ws = self.wb[L.wgp_sheet]
        out = {}
        for key, top in L.wgp_blocks.items():
            for i, seg in enumerate(("cds", "cs", "ea", "total")):
                out[f"{key}_gp_{seg}"] = _num(ws.cell(top + i, c).value)
        return out


def _num(v) -> float:
    if isinstance(v, (int, float)):
        return float(v)
    if v is None:
        return 0.0
    raise WorkbookError(f"expected a number, got {v!r}")
