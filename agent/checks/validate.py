"""Deterministic gate between Claude's output and publication (ADR-0006, ADR-0007).

    report = validate(output, source, run_context)   # recalculates a copy
    report.passed  -> publish report.recalculated_path, else fail closed

Every numeric check compares the recalculated workbook against values computed
here from the raw source (agent.checks.reference / agent.checks.metrics), never against the
workbook's own totals.

CLI:  python -m agent.checks.validate OUTPUT SOURCE [--context run_context.json] [--no-recalc]
"""
from __future__ import annotations

import argparse
import json
import re
import shutil
import tempfile
from dataclasses import asdict, dataclass, field
from datetime import date, datetime
from pathlib import Path

import openpyxl

from agent.checks.injection import describe, find_instructions
from agent.checks.metrics import PNL_LABELS, customer_gp, headline
from agent.checks.reference import (
    BALANCE_FIELDS,
    FLOW_FIELDS,
    Assumptions,
    ScheduleResult,
    cash_schedule,
    monthly_summary,
    row_labels,
    summary_labels,
)
from agent.claude.memory.prior import PriorWeek, prior_week
from agent.workbook import layout
from agent.workbook.buckets import build_buckets, month_groups
from agent.workbook.reader import TradingUpdate, _norm
from agent.workbook.recalc import recalc

EXCEL_ERRORS = ("#VALUE!", "#DIV/0!", "#REF!", "#NAME?", "#NULL!", "#NUM!", "#N/A")
SERIAL_RE = re.compile(r'(?<![\w$.:"])4[5-9]\d{3}(?![\d.])')  # Excel date serials 2023-2036 typed into formulas
TOL_EUR = 1.0  # absolute tolerance on EUR amounts
TOL_K = 0.01  # on EUR-thousand amounts


@dataclass
class Check:
    name: str
    passed: bool
    detail: str = ""
    severity: str = "error"  # "error" blocks publication, "warning" does not


@dataclass
class ValidationReport:
    checks: list[Check] = field(default_factory=list)
    recalculated_path: str | None = None
    engine: str | None = None

    @property
    def passed(self) -> bool:
        return all(c.passed for c in self.checks if c.severity == "error")

    def add(self, name: str, passed: bool, detail: str = "", severity: str = "error") -> bool:
        self.checks.append(Check(name, bool(passed), detail, severity))
        return bool(passed)

    def failures(self) -> list[Check]:
        return [c for c in self.checks if not c.passed]

    def to_dict(self) -> dict:
        return {"passed": self.passed, "engine": self.engine, "recalculated_path": self.recalculated_path,
                "checks": [asdict(c) for c in self.checks]}

    def summary(self) -> str:
        lines = [f"{'PASS' if self.passed else 'FAIL'}: {sum(c.passed for c in self.checks)}/{len(self.checks)} checks"]
        lines += [f"  [{c.severity}] {c.name}: {c.detail}" for c in self.failures()]
        return "\n".join(lines)


def company_name(context: dict) -> str:
    """The name used inside output labels: the profile's ``name``, else its key (ADR-0017)."""
    return context.get("company_name") or context.get("company") or "Inceptua"


def comparison_sheet_name(cw: int) -> str:
    return f"CW{cw:02d} vs CW{cw - 1:02d}"


# --------------------------------------------------------------------------- entry point
def validate(output: str | Path, source: str | Path, context: dict | None = None, *,
             recalculate: bool = True, engine: str | None = None, keep_dir: Path | None = None) -> ValidationReport:
    context = context or {}
    name = company_name(context)
    output, source = Path(output), Path(source)
    rep = ValidationReport()
    src = TradingUpdate(source, context.get("source_map"))  # the company's layout (ADR-0024)
    year, cw = src.iso_week
    cmp_name = comparison_sheet_name(cw)
    new_tabs = ["Cash Summary (Monthly)", "Cash Summary (Weekly)", cmp_name, "Cash Schedule"]

    if not rep.add("output file exists", output.exists(), str(output)):
        return rep
    raw = openpyxl.load_workbook(output)  # formulas, before recalculation
    if not _check_structure(rep, raw, src, new_tabs):
        return rep
    _check_source_untouched(rep, raw, source)
    _check_no_planted_instructions(rep, source)
    _check_no_hardcoded_dates(rep, raw, new_tabs)
    buckets = build_buckets(src.report_date, int(context.get("assumptions", {}).get("forward_weeks", 28)))
    _check_no_macros(rep, output)
    _check_layout(rep, raw, src, buckets, cmp_name, name)
    _check_live_formulas(rep, raw, src, buckets, cmp_name)
    if not _check_assumptions(rep, context, src):
        return rep

    # Recalculate a copy; the recalculated copy is what gets published.
    if recalculate:
        work = Path(keep_dir or tempfile.mkdtemp(prefix="vscp-validate-"))
        work.mkdir(parents=True, exist_ok=True)
        calc_path = work / output.name
        if calc_path.resolve() != output.resolve():
            shutil.copy2(output, calc_path)
        rep.engine = recalc(calc_path, engine)
    else:
        calc_path = output
    rep.recalculated_path = str(calc_path)
    vals = openpyxl.load_workbook(calc_path, data_only=True)

    _check_formula_errors(rep, vals, new_tabs)
    _check_opening_loans(rep, vals, context)
    assumptions = resolve_assumptions(context, src, vals)
    ref = cash_schedule(src.orderbook(), buckets, assumptions)
    _check_cash_schedule(rep, vals["Cash Schedule"], src, ref, name)
    _check_summary(rep, vals["Cash Summary (Weekly)"], ref, weekly=True, company=name)
    _check_summary(rep, vals["Cash Summary (Monthly)"], ref, weekly=False, company=name)

    reports_dir = context.get("prior_reports_dir")
    prior = prior_week(src, Path(reports_dir) if reports_dir else None)
    _check_comparison(rep, vals[cmp_name], src, prior)
    _check_agent_report(rep, output.parent / "agent_report.json")
    return rep


# --------------------------------------------------------------------------- checks
def _check_structure(rep, raw, src: TradingUpdate, new_tabs: list[str]) -> bool:
    names = raw.sheetnames
    missing = [t for t in new_tabs if t not in names]
    if not rep.add("generated tabs present", not missing, f"missing: {missing}" if missing else ", ".join(new_tabs)):
        return False
    rep.add("summary tabs first", names[:3] == new_tabs[:3], f"first tabs: {names[:3]}")
    before = src.layout.insert_before  # DPCache_orderbook_Data in the template; a company's map may name another
    if before and before in names:
        ok = names.index("Cash Schedule") == names.index(before) - 1
        rep.add(f"Cash Schedule before {before}", ok, f"order: {names}", "warning")
    source_names = [n for n in src.wb.sheetnames]
    kept = [n for n in names if n not in new_tabs]
    return rep.add("source tabs kept in order", kept == source_names, f"{kept} vs {source_names}")


def _check_source_untouched(rep, raw, source: Path) -> None:
    orig = openpyxl.load_workbook(source)
    changed = []
    for ws in orig.worksheets:
        out = raw[ws.title]
        for row in ws.iter_rows():
            for c in row:
                if not _same(c.value, out[c.coordinate].value):
                    changed.append(f"{ws.title}!{c.coordinate}")
                    if len(changed) >= 10:
                        break
            if len(changed) >= 10:
                break
    rep.add("source tabs unchanged", not changed, f"changed cells (first 10): {changed}" if changed else "")


def _same(a, b) -> bool:
    if isinstance(a, (int, float)) and isinstance(b, (int, float)):
        return abs(a - b) <= 1e-9 * max(1.0, abs(a))  # recalculation noise
    return a == b


INSTRUCTIONS_CHECK = "no instructions to an AI in the company's file"


def _check_no_planted_instructions(rep, source: Path) -> None:
    """A warning, never a block: the reviewer should know before approving (agent/checks/injection.py)."""
    found = find_instructions(source)
    rep.add(INSTRUCTIONS_CHECK, not found,
            "; ".join(describe(f) for f in found) + (" (the agent is told to treat file text as data; check the "
                                                     "output with this in mind)" if found else ""), "warning")


def _check_layout(rep, raw, src: TradingUpdate, buckets, cmp_name: str, company: str) -> None:
    """Same rows, columns and wording as the sample output (agent/workbook/layout.py); nothing added."""
    year, cw = src.iso_week

    def compare(name: str, expected: dict[str, object], max_row: int, max_col: int, label_col: str,
                label_rows: range, extra_ok=lambda ref: False) -> None:
        ws = raw[name]
        bad = []
        for ref, want in expected.items():
            got = ws[ref].value
            ok = want(got) if callable(want) else got == want
            if not ok:
                bad.append(f"{ref}: expected {want if not callable(want) else 'pattern'}, got {got!r}")
        for r in label_rows:  # no extra labels in the label column
            ref = f"{label_col}{r}"
            if ref not in expected and ws[ref].value not in (None, "") and not extra_ok(ref):
                bad.append(f"{ref}: unexpected {ws[ref].value!r}")
        if ws.max_row > max_row:
            bad.append(f"{ws.max_row - max_row} extra row(s) below row {max_row}")
        if ws.max_column > max_col:
            bad.append(f"{ws.max_column - max_col} extra column(s) after column {openpyxl.utils.get_column_letter(max_col)}")
        rep.add(f"'{name}' layout matches the sample", not bad, "; ".join(bad[:6]))

    starts = lambda prefix: (lambda v: isinstance(v, str) and v.startswith(prefix))  # noqa: E731
    ends = lambda suffix: (lambda v: isinstance(v, str) and v.endswith(suffix))  # noqa: E731

    # Summary tabs
    for name, labels, headers in (
            ("Cash Summary (Weekly)", layout.weekly_labels(company, cw, year, src.report_date, buckets),
             layout.weekly_headers(cw, buckets)),
            ("Cash Summary (Monthly)", layout.monthly_labels(company, cw, year, src.report_date),
             layout.monthly_headers(buckets))):
        exp = {k: (ends(" — " + v.split(" — ")[1]) if k == "A1" else v) for k, v in labels.items()}
        for j, (h4, h5) in enumerate(headers):
            col = openpyxl.utils.get_column_letter(2 + j)
            exp[f"{col}4"], exp[f"{col}5"] = h4, h5
        compare(name, exp, 25, 1 + len(headers), "A", range(1, 26))

    # Comparison tab
    exp = {"B2": layout.comparison_title(company, cw),
           "B3": starts("YTD figures in Thousands of EUR  |  "), "B5": "Metric", "C5": f"CW{cw - 1:02d}",
           "D5": f"CW{cw:02d}", "E5": "W/W Change", "F5": "W/W %", "G5": "Comment (refresh as needed)",
           "B12": lambda v: isinstance(v, str) and re.fullmatch(
               r"GP as % of FY\d{2} Budget \((€[\d.]+M|n/a)\)", v) is not None,
           "B13": starts("Adj. Revenue as % of FY"), "B34": starts("CASH & DEBT (snapshot dates "),
           "B39": starts("** Excludes Early Access pass-through revenues and COGS.")}
    for r, (_, label) in layout.COMPARISON_ROWS.items():
        if label:
            exp[f"B{r}"] = label
    names = [layout.OTHERS_LABEL.get(n, n) for n in list(customer_gp(src))[:layout.CUSTOMER_ROWS]]
    ws = raw[cmp_name]
    got_names = [ws[f"B{r}"].value for r in range(layout.CUSTOMER_FIRST_ROW, layout.CUSTOMER_FIRST_ROW + len(names))]
    rep.add("comparison lists the sample's 9 customer rows", sorted(map(str, got_names)) == sorted(names),
            f"got {got_names}")
    customer_refs = {f"B{r}" for r in range(layout.CUSTOMER_FIRST_ROW, layout.CUSTOMER_FIRST_ROW + len(names))}
    comments = [f"G{r}" for r in range(6, 38) if ws[f"G{r}"].value not in (None, "")]
    rep.add("comparison Comment column left empty", not comments, f"filled: {comments[:5]}")
    compare(cmp_name, exp, layout.FOOTNOTE_ROW, 7, "B", range(1, layout.FOOTNOTE_ROW + 1),
            extra_ok=lambda ref: ref in customer_refs)

    # Cash Schedule: header cells, totals-block labels at their offsets, the three notes lines, nothing else
    ws = raw["Cash Schedule"]
    T = 6 + len({l.so for l in src.orderbook()})
    exp = {"A1": "Calendarized Cash Flow Schedule",
           "A2": "Open Purchase EUR (negative/outflow) + Open Sales EUR (positive/inflow)",
           "A4": "Activity by Project", "B5": "Customer", "C5": "Medication", "D5": "Earliest PO Date"}
    for _, (off, label) in layout.schedule_offsets(company).items():
        if label:
            exp[f"A{T + off}"] = label
    notes_row = T + layout.SCHEDULE_OFFSETS["notes"][0]
    for n, prefix in enumerate(("- Values shown are:", "- Data source:", "- Date range:"), start=1):
        exp[f"A{notes_row + n}"] = starts(prefix)
    last_row = notes_row + layout.NOTE_LINES
    so_rows = {f"A{r}" for r in range(6, T)}
    first = _first_bucket_col(ws)
    # The last column is Margin %: catch-up + forward weeks + After (len(buckets)), then Grand Total and Margin %.
    last_col = first + len(buckets) + 1 if first else ws.max_column
    compare("Cash Schedule", exp, last_row, last_col, "A", range(1, last_row + 1), extra_ok=lambda ref: ref in so_rows)


# Totals-block rows that hold formulas in every bucket column. Only these catch-up cells may be typed
# values (SKILL.md, hard rule 1): the opening loan balance and the zero opening investment.
FORMULA_ROWS = ("outflows", "inflows", "net_operating", "loan_draws", "cash_investment", "project_investment",
                "net_inceptua", "customer_payments", "loan_repayments", "inceptua_cash", "loans_bop", "loans_repay",
                "loans_draws", "loans_eop", "bank_bop", "bank_completed", "bank_draws", "bank_purchases", "bank_repay",
                "bank_eop", "invested_bop", "invested_completed", "invested_new", "invested_eop", "total_bank",
                "total_invested", "total_cash", "total_loans", "orderbook_value", "net_debt")
TYPED_CATCHUP_ROWS = ("loans_bop", "invested_bop")


def _is_formula(v) -> bool:
    return isinstance(v, str) and v.startswith("=")


def _check_live_formulas(rep, raw, src: TradingUpdate, buckets, cmp_name: str) -> None:
    """Every calculated cell must be a live formula, not a typed number that happens to match (ADR-0006).

    Without this, an output could pass the numeric checks this week by typing in the right values,
    and silently stop working next week.
    """
    typed: dict[str, list[str]] = {}

    def need(ws, r: int, c: int) -> None:
        if not _is_formula(ws.cell(r, c).value) and ws.cell(r, c).value is not None:
            typed.setdefault(ws.title, []).append(f"{openpyxl.utils.get_column_letter(c)}{r}")
        elif ws.cell(r, c).value is None:
            typed.setdefault(ws.title, []).append(f"{openpyxl.utils.get_column_letter(c)}{r} (empty)")

    ws = raw["Cash Schedule"]
    first = _first_bucket_col(ws)
    if first is not None:
        T = 6 + len({l.so for l in src.orderbook()})
        n = len(buckets)
        daily = range(5, first - 2)  # the daily grid: E up to the two spacer columns before the catch-up bucket
        lookups = [c for c, f in ((2, "customer"), (3, "item")) if src.layout.ob_cols.get(f)]  # absent: left empty
        typed_ok = TYPED_CATCHUP_ROWS + (() if src.layout.opening_bank_cash else ("bank_bop",))
        for r in range(6, T):  # project rows: lookups, every daily cell, bucket sums, grand total, margin
            for c in (*lookups, 4, *daily, *range(first, first + n + 2)):
                need(ws, r, c)
        for key in FORMULA_ROWS:
            r = T + layout.SCHEDULE_OFFSETS[key][0]
            for i in range(n):
                if not (i == 0 and key in typed_ok):
                    need(ws, r, first + i)

    for name in ("Cash Summary (Weekly)", "Cash Summary (Monthly)"):
        ws = raw[name]
        cols = len(buckets) if "Weekly" in name else len(month_groups(buckets))
        for r in layout.SUMMARY_ROWS:
            for c in range(2, 2 + cols):
                need(ws, r, c)
        for r in layout.FLOW_ROWS:
            need(ws, r, 2 + cols)  # Grand Total

    ws = raw[cmp_name]
    absent = {k for k, v in headline(src).items() if v is None}  # not in this company's file: n/a, checked below
    metric_rows = [r for r, (key, _) in layout.COMPARISON_ROWS.items() if key and key not in absent]
    customer_rows = range(layout.CUSTOMER_FIRST_ROW, layout.CUSTOMER_FIRST_ROW + layout.CUSTOMER_ROWS)
    for r in (*metric_rows, *customer_rows):
        if ws[f"B{r}"].value not in (None, ""):
            need(ws, r, 4)  # current week: always live
            need(ws, r, 5)  # W/W change

    total = sum(len(v) for v in typed.values())
    detail = "; ".join(f"{k}: {', '.join(v[:5])}" + (f" (+{len(v) - 5} more)" if len(v) > 5 else "")
                       for k, v in typed.items())
    rep.add("calculated cells are live formulas", not typed, f"{total} typed-in or empty cells: {detail}" if typed else "")


def _check_no_macros(rep, output: Path) -> None:
    """Outputs are plain .xlsx: no VBA project, so nothing in them can ever run (ADR-0017)."""
    import zipfile
    try:
        with zipfile.ZipFile(output) as z:
            vba = [n for n in z.namelist() if n.lower().endswith("vbaproject.bin")]
    except zipfile.BadZipFile:
        vba = ["<not a valid .xlsx>"]
    rep.add("output contains no macros", not vba and output.suffix.lower() == ".xlsx",
            f"{output.name}: {vba}" if vba else output.name)


def _check_no_hardcoded_dates(rep, raw, new_tabs: list[str]) -> None:
    hits = []
    for name in new_tabs:
        for row in raw[name].iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("=") and SERIAL_RE.search(c.value):
                    hits.append(f"{name}!{c.coordinate}")
    rep.add("no hardcoded date serials in formulas", not hits, f"{len(hits)} cells, e.g. {hits[:5]}" if hits else "")


def _check_formula_errors(rep, vals, new_tabs: list[str]) -> None:
    errors = []
    for name in new_tabs:
        for row in vals[name].iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.strip() in EXCEL_ERRORS:
                    errors.append(f"{name}!{c.coordinate}={c.value}")
    rep.add("no formula errors", not errors, f"{len(errors)} errors, e.g. {errors[:5]}" if errors else "")


def _check_assumptions(rep, context: dict, src: TradingUpdate) -> bool:
    """The company's own financing terms, and a source for the opening bank cash. Never another company's."""
    a = context.get("assumptions", {})
    missing = [k for k in ("model_start_date", "po_advance_rate") if a.get(k) in (None, "")]
    if src.opening_bank_cash is None and a.get("opening_bank_cash_eur") is None:
        missing.append("opening_bank_cash_eur (the file has no opening bank cash)")
    return rep.add("company assumptions present", not missing, f"missing: {', '.join(missing)}" if missing else "")


def resolve_assumptions(context: dict, src: TradingUpdate, vals) -> Assumptions:
    a = context.get("assumptions", {})
    start = a["model_start_date"]
    loans = a.get("opening_loan_balance_eur")
    if loans is None:  # carried forward: read it from last week's output ourselves, never from the agent's
        found = prior_opening_loans(context)
        loans = found[0] if found else _catchup_value(vals["Cash Schedule"], context, "loans_bop")
    bank = src.opening_bank_cash  # the file's total bank cash (Cash!I29 in the template), else the profile's
    return Assumptions(
        model_start_date=date.fromisoformat(start) if isinstance(start, str) else start,
        po_advance_rate=float(a["po_advance_rate"]),
        opening_loan_balance=float(loans or 0.0),
        opening_bank_cash=float(bank if bank is not None else a["opening_bank_cash_eur"]),
    )


def _catchup_value(ws, context: dict, key: str):
    """A Cash Schedule catch-up cell by its row label, e.g. the opening 'Loans - BoP'."""
    row = _label_row(ws, row_labels(company_name(context))[key])
    first = _first_bucket_col(ws)
    return ws.cell(row, first).value if row and first else None


def prior_opening_loans(context: dict) -> tuple[float, str] | None:
    """The opening loan balance carried forward: 'Loans - BoP' in the catch-up column of last week's output.

    Read from the copy in the run's prior/ folder, which the pipeline checks the agent did not change. Returns
    (value, where) or None when there is no last-week output.
    """
    d = context.get("prior_reports_dir")
    if not d:
        return None
    for p in sorted(Path(d).glob("*_vClaude.xlsx")):
        try:
            ws = openpyxl.load_workbook(p, data_only=True)["Cash Schedule"]
        except (KeyError, OSError):
            continue
        value = _catchup_value(ws, context, "loans_bop")
        if isinstance(value, (int, float)):
            return float(value), f"{p.name} 'Cash Schedule' catch-up Loans - BoP"
    return None


def _check_opening_loans(rep, vals, context: dict) -> None:
    """The typed opening loan balance must equal its source: the assumption, else last week's output."""
    configured = context.get("assumptions", {}).get("opening_loan_balance_eur")
    typed = _catchup_value(vals["Cash Schedule"], context, "loans_bop")
    if configured is not None:
        expected, where = float(configured), "run_context.json assumption"
    else:
        found = prior_opening_loans(context)
        if found is None:
            rep.add("opening loan balance has a source", False,
                    "no opening_loan_balance_eur assumption and no last-week output in prior/")
            return
        expected, where = found
    ok = isinstance(typed, (int, float)) and abs(typed - expected) <= 0.01
    # The detail goes to a retry as a hint (retry_feedback), so it says where, never the value (ADR-0018).
    rep.add("opening loan balance matches its source", ok, f"the catch-up 'Loans - BoP' must equal {where}")


def _label_row(ws, label: str, col: str = "A", start: int = 1) -> int | None:
    target = _norm(label)
    for r in range(start, ws.max_row + 1):
        if _norm(ws[f"{col}{r}"].value) == target:
            return r
    return None


def _first_bucket_col(ws) -> int | None:
    """The catch-up column: the one just before the first forward week's week-ending date in row 5."""
    for c in range(5, ws.max_column + 1):
        v = ws.cell(4, c).value
        if isinstance(v, str) and re.fullmatch(r"CW\d{2} \d{4}", v.strip()):
            return c - 1
    return None


def _close(a, b, tol) -> bool:
    return isinstance(a, (int, float)) and abs(a - b) <= tol


def _check_cash_schedule(rep, ws, src: TradingUpdate, ref: ScheduleResult, company: str) -> None:
    labels = row_labels(company)
    first = _first_bucket_col(ws)
    if not rep.add("Cash Schedule bucket columns found", first is not None, "row 4 needs 'CWnn yyyy' labels"):
        return
    # Bucket headers
    bad_hdr = [b.label for i, b in enumerate(ref.buckets) if b.kind == "week"
               and not (ws.cell(4, first + i).value == b.label
                        and _as_date(ws.cell(5, first + i).value) == b.week_ending)]
    rep.add("week bucket headers match ISO weeks", not bad_hdr, f"mismatched: {bad_hdr[:5]}" if bad_hdr else "")

    # Project rows: one per unique SO, same set as the orderbook
    total_row = _label_row(ws, "Total, Net")
    sos = [ws[f"A{r}"].value for r in range(6, total_row or 6) if ws[f"A{r}"].value]
    expected = list(dict.fromkeys(l.so for l in src.orderbook()))
    rep.add("one project row per unique sales order", sorted(map(str, sos)) == sorted(expected),
            f"{len(sos)} rows vs {len(expected)} SOs")

    # Per-SO net across all buckets
    gt_col = first + len(ref.buckets)
    per_so = {}
    for l in src.orderbook():
        per_so[l.so] = per_so.get(l.so, 0.0) + (l.sales if l.so_date else 0.0) - (
            l.purchase if isinstance(l.po_date, datetime) else 0.0)
    bad_so = [so for r, so in ((r, ws[f"A{r}"].value) for r in range(6, total_row or 6)) if so
              and not _close(ws.cell(r, gt_col).value, per_so.get(str(so), 0.0), TOL_EUR)]
    rep.add("per-SO grand totals tie to orderbook", not bad_so, f"{len(bad_so)} SOs off, e.g. {bad_so[:3]}" if bad_so else "")

    # Totals block, bucket by bucket
    for fld in FLOW_FIELDS + BALANCE_FIELDS:
        row = _label_row(ws, labels[fld], start=(total_row or 1))
        if not rep.add(f"schedule row '{labels[fld]}' present", row is not None):
            continue
        off = [(ref.buckets[i].label, ws.cell(row, first + i).value, round(v, 2))
               for i, v in enumerate(ref.rows[fld]) if not _close(ws.cell(row, first + i).value or 0.0, v, TOL_EUR)]
        rep.add(f"schedule '{fld}' ties to reference", not off, f"{len(off)} buckets off, e.g. {off[:2]}" if off else "")
    for fld in ("inflows", "outflows", "net_operating"):
        row = _label_row(ws, labels[fld], start=(total_row or 1))
        if row:
            v = ws.cell(row, gt_col).value
            rep.add(f"grand total '{fld}' = orderbook", _close(v, ref.grand_total[fld], TOL_EUR),
                    f"{v} vs {ref.grand_total[fld]:.2f}")
    draws_row = _label_row(ws, labels["loan_draws"], start=(total_row or 1))
    rate = ws.cell(draws_row, 3).value if draws_row else None
    rep.add("advance rate shown next to New Loan Draws", isinstance(rate, (int, float)), f"C{draws_row}={rate!r}",
            "warning")


def _as_date(v):
    return v.date() if isinstance(v, datetime) else v


def _check_summary(rep, ws, ref: ScheduleResult, weekly: bool, company: str) -> None:
    tab = "weekly" if weekly else "monthly"
    if weekly:
        cols = len(ref.buckets)
        flows, balances = summary_labels(company)
        expected = {lbl: [v / 1000 for v in ref.rows[f]] for lbl, f in {**flows, **balances}.items()}
    else:
        labels, expected = monthly_summary(ref, company)
        cols = len(labels)
    for label, values in expected.items():
        row = _label_row(ws, label)
        if not rep.add(f"{tab} summary row '{label}' present", row is not None):
            continue
        off = [(i, ws.cell(row, 2 + i).value, round(v, 3)) for i, v in enumerate(values[:cols])
               if not _close(ws.cell(row, 2 + i).value or 0.0, v, TOL_K)]
        rep.add(f"{tab} summary '{label}' ties to schedule", not off, f"{len(off)} columns off, e.g. {off[:2]}" if off else "")


def _find_metric_row(ws, label: str, start: int = 1, exact: bool = False) -> int | None:
    target = _norm(label)
    for r in range(start, ws.max_row + 1):
        v = _norm(ws[f"B{r}"].value)
        if v and (v == target or (not exact and v.startswith(target))):
            return r
    return None


def _is_blank_or_na(v) -> bool:
    return v is None or (isinstance(v, str) and v.strip().lower() in ("", "n/a", "na", "–", "-"))


def _check_comparison(rep, ws, src: TradingUpdate, prior: PriorWeek) -> None:
    cur = headline(src)
    rep.add("prior-week source identified", prior.source != "none", prior.source, "warning")
    for key, label in PNL_LABELS.items():
        row = _find_metric_row(ws, label)
        if not rep.add(f"comparison row '{label}' present", row is not None):
            continue
        tol = 1e-4 if key.endswith("_pct") else 0.5
        c_val, p_val = ws[f"D{row}"].value, ws[f"C{row}"].value
        if cur[key] is None:  # the company's file has no such input: n/a both weeks, never a number
            rep.add(f"'{key}' shown as n/a (not in this company's file)", _is_blank_or_na(c_val) and _is_blank_or_na(p_val),
                    f"got {c_val!r} / {p_val!r}")
            continue
        rep.add(f"current '{key}' matches source", _close(c_val, cur[key], tol), f"{c_val} vs {cur[key]:.4f}")
        if key in prior.values:
            rep.add(f"prior '{key}' matches {prior.source}", _close(p_val, prior.values[key], tol),
                    f"{p_val} vs {prior.values[key]:.4f}")
        else:
            rep.add(f"prior '{key}' not invented", _is_blank_or_na(p_val), f"expected n/a, got {p_val!r}")
    oo = _find_metric_row(ws, PNL_LABELS["open_orders"])
    if oo:
        v = float(ws[f"D{oo}"].value or 0)
        rep.add("open orders is a whole number", abs(v - round(v)) < 1e-6, str(ws[f"D{oo}"].value))

    # Customer moves: the first 9 CDS customer rows on Customer-Level (as in the sample), matched by name
    customers = dict(list(customer_gp(src).items())[:layout.CUSTOMER_ROWS])
    start = _find_metric_row(ws, "KEY CUSTOMER MOVES") or 1
    missing, wrong = [], []
    for name, gp in customers.items():
        row = _find_metric_row(ws, layout.OTHERS_LABEL.get(name, name), start, exact=True)
        if row is None:
            missing.append(name)
            continue
        if not _close(ws[f"D{row}"].value, gp, 0.5):
            wrong.append(f"{name}: {ws[f'D{row}'].value} vs {gp:.1f}")
        if prior.customers:
            if not _close(ws[f"C{row}"].value or 0.0, prior.customers.get(name, 0.0), 0.5):
                wrong.append(f"{name} prior: {ws[f'C{row}'].value} vs {prior.customers.get(name, 0.0):.1f}")
        elif not _is_blank_or_na(ws[f"C{row}"].value):
            wrong.append(f"{name} prior invented: {ws[f'C{row}'].value}")
    rep.add("customer rows present (matched by name)", not missing, f"missing: {missing}" if missing else "")
    rep.add("customer values match source and prior week", not wrong, "; ".join(wrong[:4]))


def _check_agent_report(rep, path: Path) -> None:
    if not rep.add("agent_report.json written", path.exists(), str(path), "warning"):
        return
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as e:
        rep.add("agent_report.json is valid JSON", False, str(e), "warning")
        return
    rep.add("agent_report.json lists value sources", bool(data.get("sources") or data.get("provenance")), "", "warning")


# --------------------------------------------------------------------------- retry feedback
BUCKET_RE = re.compile(r"After CW\d{2} \d{4}|CW\d{2} \(through [^)]*\)|CW\d{2} \d{4}")
NUMERIC_CHECKS = ("ties to", "matches source", "matches prior", "matches in_file", "= orderbook", "per-SO grand totals",
                  "customer values match", "whole number")


def retry_feedback(report: dict | ValidationReport) -> list[str]:
    """Failed checks as hints for the next attempt: what and where, never the expected values.

    Expected numbers are withheld on purpose. Given them, an agent could type them in to pass; the
    fix must be in the formula (and ``calculated cells are live formulas`` catches typed numbers).
    """
    data = report.to_dict() if isinstance(report, ValidationReport) else report
    hints = []
    for c in data.get("checks", []):
        if c.get("passed") or c.get("severity") != "error":
            continue
        name, detail = c["name"], str(c.get("detail", ""))
        if any(k in name for k in NUMERIC_CHECKS):
            where = list(dict.fromkeys(BUCKET_RE.findall(detail)))
            if "customer" in name:
                where = list(dict.fromkeys(m.strip() for m in re.findall(r"(?:^|; )([^:;]+?):", detail)))
            hint = f"{name}: FAILED (wrong values"
            hint += f" in {', '.join(where[:8])}" + (" and more" if len(where) > 8 else "") + ")" if where else ")"
        else:
            hint = f"{name}: FAILED ({detail[:300]})" if detail else f"{name}: FAILED"
        hints.append(hint)
    return hints


def main() -> None:
    p = argparse.ArgumentParser(description="Validate a _vClaude output workbook")
    p.add_argument("output", type=Path)
    p.add_argument("source", type=Path)
    p.add_argument("--context", type=Path, help="run_context.json")
    p.add_argument("--no-recalc", action="store_true", help="trust cached values (e.g. a file saved by Excel)")
    p.add_argument("--json", type=Path, help="write the report here")
    a = p.parse_args()
    ctx = json.loads(a.context.read_text(encoding="utf-8")) if a.context else {}
    rep = validate(a.output, a.source, ctx, recalculate=not a.no_recalc)
    print(rep.summary())
    if a.json:
        a.json.write_text(json.dumps(rep.to_dict(), indent=2, default=str), encoding="utf-8")
    raise SystemExit(0 if rep.passed else 1)


if __name__ == "__main__":
    main()
